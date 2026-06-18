"""
PPTX Review Tool
Renders each slide as a high-res PNG using PowerPoint COM automation,
and extracts structured text content via python-pptx.
Usage: python pptx_review.py <path_to_pptx> [--out-dir <dir>]
"""
import sys
import os
import json
import argparse
from pathlib import Path

# Force UTF-8 output on Windows
if sys.platform == "win32":
    import io
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
    sys.stderr = io.TextIOWrapper(sys.stderr.buffer, encoding="utf-8", errors="replace")

def render_slides(pptx_path: str, out_dir: str, width: int = 1920, height: int = 1080) -> list[str]:
    """Export each slide as PNG using PowerPoint COM. Returns list of PNG paths."""
    import win32com.client
    import pythoncom

    pythoncom.CoInitialize()
    pptx_abs = str(Path(pptx_path).resolve())
    out_abs = str(Path(out_dir).resolve())
    os.makedirs(out_abs, exist_ok=True)

    ppt = win32com.client.Dispatch("PowerPoint.Application")
    ppt.Visible = True

    try:
        prs = ppt.Presentations.Open(pptx_abs, ReadOnly=True, Untitled=False, WithWindow=False)
        paths = []
        for i in range(1, prs.Slides.Count + 1):
            slide = prs.Slides(i)
            out_path = os.path.join(out_abs, f"slide_{i:02d}.png")
            slide.Export(out_path, "PNG", width, height)
            paths.append(out_path)
            print(f"  Rendered slide {i}/{prs.Slides.Count} -> {out_path}", file=sys.stderr)
        prs.Close()
        return paths
    finally:
        ppt.Quit()
        pythoncom.CoUninitialize()


def extract_text(pptx_path: str) -> list[dict]:
    """Extract structured text from each slide using python-pptx."""
    from pptx import Presentation
    from pptx.util import Pt

    prs = Presentation(pptx_path)
    slides = []
    for i, slide in enumerate(prs.slides, 1):
        title = ""
        body_texts = []
        notes = ""

        # Title
        if slide.shapes.title and slide.shapes.title.has_text_frame:
            title = slide.shapes.title.text.strip()

        # All text frames
        for shape in slide.shapes:
            if shape == slide.shapes.title:
                continue
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    text = para.text.strip()
                    if text:
                        body_texts.append(text)
            try:
                from pptx.enum.shapes import PP_PLACEHOLDER
                from pptx.util import Emu
                if shape.shape_type == 19:  # MSO_SHAPE_TYPE.TABLE
                    rows = []
                    for row in shape.table.rows:
                        rows.append([cell.text.strip() for cell in row.cells])
                    body_texts.append({"table": rows})
            except Exception:
                pass

        # Speaker notes
        if slide.has_notes_slide:
            notes_tf = slide.notes_slide.notes_text_frame
            notes = notes_tf.text.strip() if notes_tf else ""

        slides.append({
            "slide": i,
            "title": title,
            "content": body_texts,
            "notes": notes,
        })
    return slides


def main():
    parser = argparse.ArgumentParser(description="Render and extract PPTX content")
    parser.add_argument("pptx", help="Path to .pptx file")
    parser.add_argument("--out-dir", default=None, help="Output directory for PNGs (default: <pptx_dir>/slides_out)")
    parser.add_argument("--no-render", action="store_true", help="Skip PNG rendering, text only")
    parser.add_argument("--json", action="store_true", help="Output text as JSON")
    args = parser.parse_args()

    pptx_path = args.pptx
    if not os.path.exists(pptx_path):
        print(f"ERROR: file not found: {pptx_path}", file=sys.stderr)
        sys.exit(1)

    out_dir = args.out_dir or str(Path(pptx_path).parent / "slides_out")

    print(f"Extracting text content...", file=sys.stderr)
    slides = extract_text(pptx_path)

    if not args.no_render:
        print(f"Rendering slides to {out_dir}...", file=sys.stderr)
        png_paths = render_slides(pptx_path, out_dir)
        for slide, png in zip(slides, png_paths):
            slide["png"] = png

    if args.json:
        print(json.dumps(slides, indent=2))
    else:
        for s in slides:
            print(f"\n=== Slide {s['slide']}: {s['title']} ===")
            for item in s["content"]:
                if isinstance(item, dict) and "table" in item:
                    print("  [TABLE]")
                    for row in item["table"]:
                        print("   | " + " | ".join(row))
                else:
                    print(f"  {item}")
            if s["notes"]:
                print(f"  [NOTES] {s['notes'][:200]}")
            if "png" in s:
                print(f"  [PNG] {s['png']}")


if __name__ == "__main__":
    main()
